#!/usr/bin/env python3
"""Extract auditable R-01 sample metadata from the bundled atlas Table S2.

This extractor reads only workbook metadata. It does not inspect expression,
image, archive, or derived annotation assets. Reported TLS presence/count are
retained as potential provenance and are never accepted as ground truth here.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


SHEET_NAME = "Table S2"
HEADER_ROW = 3
EXPECTED_HEADERS = [
    "Cancer",
    "Sample size",
    "Patients",
    "Sample ID",
    "Tumor stage",
    "Histological subtype",
    "ST/Visium",
    "FFPE/Fresh Frozen",
    "TLS presence",
    "TLS count",
    "Reference",
    "Data availability",
]
SAMPLE_FIELDS = [
    "record_id",
    "source_asset",
    "source_sheet",
    "source_row",
    "cancer_raw",
    "study_namespace",
    "patient_id_raw",
    "sample_id_raw",
    "tumor_stage_raw",
    "histological_subtype_raw",
    "platform_raw",
    "preservation_raw",
    "reference_raw",
    "data_availability_raw",
    "tls_presence_raw",
    "tls_count_raw",
    "potential_gt_provenance",
    "gt_acceptance_status",
]
EVIDENCE_FIELDS = [
    "evidence_id",
    "record_id",
    "field",
    "source_asset",
    "source_sheet",
    "source_range",
    "raw_value",
    "evidence_grade",
    "interpretation",
    "claim_eligibility",
]
SOURCE_FIELD_MAP = [
    ("cancer_raw", 1),
    ("patient_id_raw", 3),
    ("sample_id_raw", 4),
    ("tumor_stage_raw", 5),
    ("histological_subtype_raw", 6),
    ("platform_raw", 7),
    ("preservation_raw", 8),
    ("tls_presence_raw", 9),
    ("tls_count_raw", 10),
    ("reference_raw", 11),
    ("data_availability_raw", 12),
]
GT_FIELDS = {"tls_presence_raw", "tls_count_raw"}
SOURCE_ASSET = "paper/tables/science.adz2742_tables_s1_to_s8.xlsx"


@dataclass(frozen=True)
class SourcedValue:
    value: Any
    source_range: str


def _raw_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _normalized_text(value: Any) -> str:
    return " ".join(_raw_text(value).replace("\xa0", " ").split())


def _is_unknown_metadata(value: Any) -> bool:
    return _normalized_text(value).casefold() in {
        "",
        "-",
        "na",
        "n/a",
        "not available",
        "not applicable",
        "unknown",
    }


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _merged_value_map(sheet: Any) -> dict[tuple[int, int], SourcedValue]:
    resolved: dict[tuple[int, int], SourcedValue] = {}
    for merged_range in sheet.merged_cells.ranges:
        min_col, min_row, max_col, max_row = merged_range.bounds
        value = sheet.cell(min_row, min_col).value
        source_range = str(merged_range)
        for row in range(min_row, max_row + 1):
            for column in range(min_col, max_col + 1):
                resolved[(row, column)] = SourcedValue(value, source_range)
    return resolved


def _value_at(
    sheet: Any,
    merged_values: dict[tuple[int, int], SourcedValue],
    row: int,
    column: int,
) -> SourcedValue:
    if (row, column) in merged_values:
        return merged_values[(row, column)]
    cell = sheet.cell(row, column)
    return SourcedValue(cell.value, cell.coordinate)


def _study_namespace(
    cancer: SourcedValue,
    reference: SourcedValue,
    availability: SourcedValue,
) -> str:
    normalized_cancer = _normalized_text(cancer.value) or "unknown-cancer"
    readable_cancer = re.sub(r"[^A-Za-z0-9._-]+", "-", normalized_cancer).strip("-").lower()
    identity = {
        "cancer": normalized_cancer,
        "reference": _normalized_text(reference.value),
        "data_availability": _normalized_text(availability.value),
        "reference_source_range": reference.source_range,
        "data_availability_source_range": availability.source_range,
    }
    digest = hashlib.sha256(
        json.dumps(identity, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()[:12]
    return f"atlas_s2::{readable_cancer or 'unknown-cancer'}::{digest}"


def _is_total_row(sample_size: SourcedValue) -> bool:
    return _normalized_text(sample_size.value).casefold().rstrip(":") == "total"


def _extract(
    workbook_path: Path,
) -> tuple[list[dict[str, str]], list[dict[str, str]], dict[str, int]]:
    workbook_path = workbook_path.resolve()
    workbook = load_workbook(workbook_path, read_only=False, data_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"missing required sheet: {SHEET_NAME}")
    sheet = workbook[SHEET_NAME]
    headers = [_raw_text(sheet.cell(HEADER_ROW, column).value) for column in range(1, 13)]
    if headers != EXPECTED_HEADERS:
        raise ValueError(f"unexpected {SHEET_NAME} header: {headers!r}")

    merged_values = _merged_value_map(sheet)
    records: list[dict[str, str]] = []
    evidence: list[dict[str, str]] = []
    excluded_total_rows = 0

    for row in range(HEADER_ROW + 1, sheet.max_row + 1):
        sample_size = _value_at(sheet, merged_values, row, 2)
        if _is_total_row(sample_size):
            excluded_total_rows += 1
            continue
        sample = _value_at(sheet, merged_values, row, 4)
        if not _normalized_text(sample.value):
            continue

        sourced = {
            field: _value_at(sheet, merged_values, row, column)
            for field, column in SOURCE_FIELD_MAP
        }
        record_id = f"atlas_s2_row_{row:04d}"
        record = {
            "record_id": record_id,
            "source_asset": SOURCE_ASSET,
            "source_sheet": SHEET_NAME,
            "source_row": str(row),
            **{field: _raw_text(value.value) for field, value in sourced.items()},
            "study_namespace": _study_namespace(
                sourced["cancer_raw"],
                sourced["reference_raw"],
                sourced["data_availability_raw"],
            ),
            "potential_gt_provenance": f"{SHEET_NAME}!I{row}:J{row}; reported summary only",
            "gt_acceptance_status": "NOT_ACCEPTED_R01_METADATA_ONLY",
        }
        records.append({field: record[field] for field in SAMPLE_FIELDS})

        for field_index, (field, _) in enumerate(SOURCE_FIELD_MAP, 1):
            value = sourced[field]
            is_gt_field = field in GT_FIELDS
            is_unknown = _is_unknown_metadata(value.value)
            evidence.append(
                {
                    "evidence_id": f"atlas_s2_r{row:04d}_f{field_index:02d}",
                    "record_id": record_id,
                    "field": field,
                    "source_asset": SOURCE_ASSET,
                    "source_sheet": SHEET_NAME,
                    "source_range": value.source_range,
                    "raw_value": _raw_text(value.value),
                    "evidence_grade": "E0_unknown" if is_unknown else "E3_explicit",
                    "interpretation": (
                        "No substantive value is provided by the bundled Table S2 metadata."
                        if is_unknown
                        else "Reported TLS summary; provenance candidate only, not independently accepted GT."
                        if is_gt_field
                        else "Direct value from bundled Table S2 metadata."
                    ),
                    "claim_eligibility": (
                        "POTENTIAL_GT_ONLY_NOT_ACCEPTED"
                        if is_gt_field
                        else "R01_IDENTITY_METADATA_ONLY"
                    ),
                }
            )

    stats = {
        "sample_record_count": len(records),
        "evidence_record_count": len(evidence),
        "excluded_total_row_count": excluded_total_rows,
    }
    return records, evidence, stats


def extract_atlas_records(
    workbook_path: Path | str,
) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    records, evidence, _ = _extract(Path(workbook_path))
    return records, evidence


def _write_tsv(path: Path, fieldnames: Iterable[str], rows: Iterable[dict[str, str]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=list(fieldnames),
            delimiter="\t",
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerows(rows)


def write_atlas_outputs(
    workbook_path: Path | str,
    output_dir: Path | str,
) -> dict[str, Any]:
    workbook_path = Path(workbook_path).resolve()
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    records, evidence, stats = _extract(workbook_path)

    samples_path = output_dir / "atlas_samples.tsv"
    evidence_path = output_dir / "atlas_evidence.tsv"
    manifest_path = output_dir / "atlas_extraction_manifest.json"
    _write_tsv(samples_path, SAMPLE_FIELDS, records)
    _write_tsv(evidence_path, EVIDENCE_FIELDS, evidence)

    manifest = {
        "extractor": "scripts/r01_extract_atlas.py",
        "source_asset": SOURCE_ASSET,
        "source_sha256": _sha256_file(workbook_path),
        "sheets_read": [SHEET_NAME],
        **stats,
        "merged_cell_policy": "Use the value and source range of the containing merged cell; do not forward-fill ordinary blanks.",
        "total_row_policy": "Exclude rows whose Sample size metadata is Total or Total:.",
        "ground_truth_policy": "TLS fields are potential provenance only; no GT is accepted by R-01.",
        "unknown_geometry_policy": "Do not create or infer block, z-position, section order, spacing, or thickness fields.",
        "outputs": [samples_path.name, evidence_path.name],
    }
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    return manifest


def main() -> None:
    project_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--workbook",
        type=Path,
        default=project_root / SOURCE_ASSET,
        help="bundled atlas supplementary workbook",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=project_root / "infra/sample-registry/staging",
        help="directory for atlas_* audit outputs",
    )
    args = parser.parse_args()
    summary = write_atlas_outputs(args.workbook, args.output_dir)
    print(json.dumps(summary, ensure_ascii=False, sort_keys=True))


if __name__ == "__main__":
    main()
