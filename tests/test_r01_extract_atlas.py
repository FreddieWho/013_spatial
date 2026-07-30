import csv
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from scripts.r01_extract_atlas import extract_atlas_records, write_atlas_outputs


HEADERS = [
    "Cancer",
    "Sample size",
    "Patients",
    "Sample ID",
    "Tumor stage",
    "Histological subtype",
    "ST/Visium",
    "FFPE/Fresh Frozen",
    "TLS presence",
    "TLS count",
    "Reference",
    "Data availability",
]


def make_fixture(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Table S2"
    sheet.append(["Table S2 fixture"])
    sheet.append([])
    sheet.append(HEADERS)
    sheet.append(
        [
            "LUAD",
            2,
            "P-01",
            "S-01",
            "III",
            "Adenocarcinoma",
            "Visium",
            "FFPE",
            "Yes",
            3,
            " PMID: 123 ",
            " GSE000001 ",
        ]
    )
    sheet.append(
        [
            None,
            None,
            None,
            "S-02",
            None,
            "Adenocarcinoma",
            "Visium",
            "Fresh Frozen",
            "No",
            0,
            None,
            None,
        ]
    )
    sheet.append(
        [
            "KIRC",
            1,
            "P-02",
            "S-03",
            "-",
            "clear cell renal cell carcinoma",
            "ST",
            "NA",
            "Yes",
            1,
            "In-house",
            "GSE000002",
        ]
    )
    sheet.append([None, "Total:", 2, 3, None, None, None, None, 2, 4])
    sheet.merge_cells("A4:A5")
    sheet.merge_cells("B4:B5")
    sheet.merge_cells("C4:C5")
    sheet.merge_cells("K4:K5")
    sheet.merge_cells("L4:L5")
    workbook.save(path)


class ExtractAtlasTests(unittest.TestCase):
    def test_extracts_sample_rows_and_resolves_only_merged_metadata(self):
        with tempfile.TemporaryDirectory() as tmp:
            workbook = Path(tmp) / "atlas.xlsx"
            make_fixture(workbook)

            records, evidence = extract_atlas_records(workbook)

            self.assertEqual([record["sample_id_raw"] for record in records], ["S-01", "S-02", "S-03"])
            self.assertEqual(records[1]["cancer_raw"], "LUAD")
            self.assertEqual(records[1]["patient_id_raw"], "P-01")
            self.assertEqual(records[1]["tumor_stage_raw"], "")
            self.assertEqual(records[1]["reference_raw"], " PMID: 123 ")
            self.assertEqual(records[1]["data_availability_raw"], " GSE000001 ")
            self.assertEqual(records[0]["study_namespace"], records[1]["study_namespace"])
            self.assertNotEqual(records[1]["study_namespace"], records[2]["study_namespace"])
            self.assertEqual(len(evidence), 3 * 11)
            patient_evidence = next(
                item
                for item in evidence
                if item["record_id"] == records[1]["record_id"]
                and item["field"] == "patient_id_raw"
            )
            self.assertEqual(patient_evidence["source_range"], "C4:C5")
            stage_evidence = next(
                item
                for item in evidence
                if item["record_id"] == records[1]["record_id"]
                and item["field"] == "tumor_stage_raw"
            )
            self.assertEqual(stage_evidence["source_range"], "E5")
            self.assertEqual(stage_evidence["evidence_grade"], "E0_unknown")

    def test_tls_fields_are_potential_provenance_not_accepted_ground_truth(self):
        with tempfile.TemporaryDirectory() as tmp:
            workbook = Path(tmp) / "atlas.xlsx"
            make_fixture(workbook)

            records, evidence = extract_atlas_records(workbook)

            self.assertEqual(records[0]["tls_presence_raw"], "Yes")
            self.assertEqual(records[0]["tls_count_raw"], "3")
            self.assertEqual(records[0]["gt_acceptance_status"], "NOT_ACCEPTED_R01_METADATA_ONLY")
            self.assertIn("Table S2!I4:J4", records[0]["potential_gt_provenance"])
            forbidden = {"gt_label", "accepted_gt", "block_id", "z_position", "section_order"}
            self.assertTrue(forbidden.isdisjoint(records[0]))
            tls_evidence = [
                item for item in evidence if item["field"] in {"tls_presence_raw", "tls_count_raw"}
            ]
            self.assertTrue(tls_evidence)
            self.assertTrue(
                all(item["claim_eligibility"] == "POTENTIAL_GT_ONLY_NOT_ACCEPTED" for item in tls_evidence)
            )

    def test_outputs_are_deterministic_and_manifest_records_policy(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook = root / "atlas.xlsx"
            output = root / "staging"
            make_fixture(workbook)

            summary1 = write_atlas_outputs(workbook, output)
            first = {path.name: path.read_bytes() for path in output.glob("atlas_*")}
            summary2 = write_atlas_outputs(workbook, output)
            second = {path.name: path.read_bytes() for path in output.glob("atlas_*")}

            self.assertEqual(summary1, summary2)
            self.assertEqual(first, second)
            with (output / "atlas_samples.tsv").open(newline="", encoding="utf-8") as handle:
                rows = list(csv.DictReader(handle, delimiter="\t"))
            self.assertEqual(len(rows), 3)
            self.assertNotIn("block_id", rows[0])
            self.assertNotIn("z_position", rows[0])
            manifest = json.loads((output / "atlas_extraction_manifest.json").read_text())
            self.assertEqual(manifest["sample_record_count"], 3)
            self.assertEqual(manifest["excluded_total_row_count"], 1)
            self.assertEqual(manifest["ground_truth_policy"], "TLS fields are potential provenance only; no GT is accepted by R-01.")
            self.assertEqual(manifest["sheets_read"], ["Table S2"])


if __name__ == "__main__":
    unittest.main()
